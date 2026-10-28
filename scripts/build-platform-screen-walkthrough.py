"""Build a short editable handbook from approved actual-component captures."""
from pathlib import Path
import json
import hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'docs/storybook'
PAGES = [
 ('household-request', '1 Household request',
  'An adult adds photos, an address and a collection time. Choose pickup or drop off and verify the phone. The household does not need to guess the material weight.',
  'The shop must accept the request. The household can then follow its collection record.'),
 ('household-receipt', '2 Receipt and payment review',
  'The shop records the measured material, weight and agreed amount. The household can accept the receipt or report a wrong line. The original record stays in the history.',
  'Receipt approval and payment acknowledgement are separate. The shop pays the household directly; Luma records what each party reports.'),
 ('shop-workspace', '3 Daily shop work',
  'The workspace puts new requests and planned collections first. The team can open stock, prices, buying and selling records. Access depends on each person’s role.',
  'Stock records keep exact quantities. Recorded cost is separate from offer value. A business can review the day’s stock and choose its own low stock alert. An assigned worker records the measured amount collected. The shop confirms the amount received; a difference needs a reason. That handover does not add stock again or pay anyone.'),
 ('business-material', '4 Material and processing records',
  'A business records material, its measured quantity and physical hand offs. Processing links each input to measured outputs, residuals and loss. Unknown costs stay unknown.',
  'A handover note alone does not add stock, confirm payment or prove legal ownership. Processing, tests, sale terms and permitted handling have their own checks.'),
 ('services', '5 Services and programmes',
  'A person or business can make a private service enquiry. The chosen model distinguishes a material sale from paid collection, processing or treatment. Programme records can describe an area, period and stream.',
  'A chosen recipient sees only the shared enquiry or report. A record does not award a loan, benefit, donation payment or official certificate.'),
 ('admin', '6 Admin work',
  'The admin reviews records, follows support requests and helps resolve problems. Admin sign in has extra protection. Moving a collection offer still requires an eligible shop and its acceptance.',
  'Changes keep a history. A review or correction does not manufacture a permit, provider approval or payment.'),
 ('worker-handover', '7 Check the worker handover',
  'The assigned worker records the measured amount collected and its evidence reference. After the shop completes its measured receipt, the shop confirms the amount received. A difference needs a reason.',
  'The handover links existing records. It does not add stock again, settle wages, pay the household or give the worker access to the household’s private details.'),
 ('account-removal', '8 Request account removal',
  'The account owner checks what removal means, signs in again when needed, and confirms the request. Business ownership or unfinished duties can hold removal. A held request leaves the account open.',
  'Eligible removal closes access and sharing. Required historical evidence stays for a separate review. Removal does not settle money or erase every record.'),
]

def picture_detail(document, path, detail):
    from PIL import Image
    with Image.open(path) as image:
        iw, ih = image.size
    crop = detail or {'left': 0, 'top': 0, 'width': iw, 'height': ih}
    left, top, width, height = [crop[k] for k in ('left', 'top', 'width', 'height')]
    assert 0 <= left < left + width <= iw and 0 <= top < top + height <= ih
    run = document.add_paragraph().add_run()
    display_width = min(6.3, 5.1 * width / height)
    shape = run.add_picture(str(path), width=Inches(display_width))
    shape.height = Inches(display_width * height / width)
    src = OxmlElement('a:srcRect')
    for key, value in [('l', left / iw), ('t', top / ih), ('r', (iw-left-width)/iw), ('b', (ih-top-height)/ih)]:
        src.set(key, str(round(value * 100000)))
    shape._inline.graphic.graphicData.pic.blipFill.insert(1, src)


def build():
    manifest = json.loads((BASE / 'screenshots.json').read_text())
    if manifest['status'] != 'verified_current_release':
        raise ValueError('Current reviewed screenshot evidence is required')
    capture = ROOT / manifest['captureManifest']
    assert hashlib.sha256(capture.read_bytes()).hexdigest() == manifest['captureManifestSha256']
    capture_evidence = json.loads(capture.read_text())
    if capture_evidence.get("status") != "approved-isolated-current-components-subset":
        raise ValueError("An approved selected-image receipt is required")
    approved = {entry["file"]: entry for entry in capture_evidence["captures"]}
    for screen in manifest["screens"].values():
        entry = approved.get(Path(screen["path"]).name)
        if not entry or entry.get("visualReview") != "passed" or entry["screenshotSha256"] != screen["sha256"]:
            raise ValueError("A selected screenshot lacks exact visual approval")
    for source, digest in capture_evidence['sourceHashes'].items():
        assert hashlib.sha256((ROOT / source).read_bytes()).hexdigest() == digest, source
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.65), Inches(.65)
    section.left_margin, section.right_margin = Inches(1.1), Inches(1.1)
    for key in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        style = d.styles[key]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12 if key == 'Normal' else 11 if key == 'Caption' else 22 if key == 'Heading 1' else 30 if key == 'Title' else 12)
        style.paragraph_format.space_after = Pt(10)
    # The blank Word template carries blue title borders; remove inherited decoration.
    for style in d.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    d.styles['Normal'].paragraph_format.line_spacing = 1.12
    d.styles['Title'].paragraph_format.space_after = Pt(18)
    header = section.header.paragraphs[0]
    header.text = 'Luma Green web platform examples'
    header.style = d.styles['Caption']
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    d.add_paragraph('Luma Green screen walkthrough', 'Title')
    d.add_paragraph('Web platform examples 28 October 2026', 'Subtitle')
    d.add_paragraph('This short guide explains eight key screens and how their records connect. Use it with the complete platform plan. The pictures show current app components in an isolated test with made up data. They do not prove a real account, sign in, deployment or provider transaction.')
    d.add_heading('Follow the material', level=1)
    for text in ['A household asks for pickup or drop off.', 'A shop accepts, weighs and records the receipt.', 'The household reviews the receipt and acknowledges payment separately.', 'Businesses record stock, trade, processing and physical hand offs.', 'Partners see only explicitly shared work or reports.', 'Admin and support help with recorded problems.']:
        d.add_paragraph(text, 'List Number')
    d.add_heading('Keep the checks separate', level=1)
    d.add_paragraph('A material movement is not a payment. A programme record is not an award. A reference definition is not permission to handle specialist material. Missing services must show that they are unavailable.')
    for key, title, text, action in PAGES:
        d.add_page_break()
        d.add_heading(title, level=1)
        d.add_paragraph(text)
        d.add_paragraph(action)
        screen = manifest['screens'][key]
        path = ROOT / screen['path']
        assert hashlib.sha256(path.read_bytes()).hexdigest() == screen['sha256']
        picture_detail(d, path, screen.get('detail'))
        d.add_paragraph(screen['caption'], 'Caption')
    d.add_page_break()
    d.add_heading('What the examples prove', level=1)
    d.add_paragraph('The screenshots show how the current components look with controlled example data. The full original captures are preserved with the capture evidence. Detail views select part of those originals; they do not change the app or invent a screen.')
    d.add_paragraph('The characters and pictures in the separate storybook are fictional. App screenshots and illustrations have different purposes. Read the caption before treating a picture as evidence.')
    d.add_heading('What needs real setup', level=1)
    d.add_paragraph('Messages, business verification, payments, private storage and long running jobs depend on the selected services and their configuration. Account approval and a successful real transaction are separate checks. Local examples do not prove them.')
    d.add_paragraph('Analytics and error reporting stay off until configured. Marketing analytics needs the visitor’s choice. Private forms and records must not be added to tracking events.')
    d.add_heading('Current scope', level=1)
    d.add_paragraph('This guide covers the web platform. Expo iOS and Android work is paused. Existing app work is saved for a later decision.')
    out = ROOT / 'output/docx/luma-green-screen-walkthrough.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    print(out)

if __name__ == '__main__':
    build()
